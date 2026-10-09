"""文档加载 —— 按扩展名分派 Loader（PDF / DOCX / TXT / MD）。

PDF 走 PyMuPDF，DOCX 走 python-docx，两者都是可选依赖：没装也能跑纯文本流程。
"""

from __future__ import annotations

from pathlib import Path


def load_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    return path.read_text(encoding="utf-8", errors="ignore")


def load_bytes(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _load_pdf_bytes(data)
    if suffix == ".docx":
        return _load_docx_bytes(data)
    return data.decode("utf-8", errors="ignore")


def _load_pdf(path: Path) -> str:
    try:
        import pymupdf  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析 PDF 需要安装 pymupdf：pip install pymupdf") from exc
    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _load_pdf_bytes(data: bytes) -> str:
    try:
        import pymupdf  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析 PDF 需要安装 pymupdf：pip install pymupdf") from exc
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def _load_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析 DOCX 需要安装 python-docx：pip install python-docx") from exc
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _load_docx_bytes(data: bytes) -> str:
    import io

    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析 DOCX 需要安装 python-docx：pip install python-docx") from exc
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)
